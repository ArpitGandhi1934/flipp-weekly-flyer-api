"""Shared by search.py, flyers.py and weekly.py: the Actor call, the flat output row, sorting, and
the CSV / Excel (.xlsx) / JSONL writers. flyerlib.mjs is the Node.js twin; both write
byte-identical CSV and JSONL."""

from __future__ import annotations

import csv
import json
import os
import re
import sys
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

ACTOR = "yugenox/flipp-flyer-deals-scraper"

# From the Actor's pricing in the Apify API, read 2026-10-06. The Store page shows the current price.
PRICE_PER_DEAL = 0.0015  # every deal saved
PRICE_PER_DETAILS = 0.0005  # add-on per deal with --details
START_FEE = 0.001  # per run (per GB of memory, 1 GB by default) ...
START_FEE_FROM = "2026-10-11T19:57:49+00:00"  # ... from this moment; $0 before it

# Column order for CSV, XLSX and JSONL.
COLUMNS = [
    "search_term", "rank", "price_basis", "name", "brand", "merchant", "effective_price",
    "price_per_kg", "price_unit", "price_as_printed", "sale_story", "current_price", "original_price",
    "discount_pct", "dollars_off", "multi_buy_qty", "multi_buy_total", "currency", "valid_from", "valid_to",
    "item_type", "store_address", "description", "sku", "product_url", "image_url", "flyer_id", "flyer_name",
    "flyer_categories", "product_category", "in_store_only", "rating", "review_count", "postal_code",
    "change_type", "previous_price", "item_id", "merchant_id", "scraped_at",
]
NUMBER_COLUMNS = {
    "rank", "effective_price", "price_per_kg", "current_price", "original_price", "discount_pct",
    "dollars_off", "multi_buy_qty", "multi_buy_total", "rating", "review_count", "previous_price",
}
DATE_COLUMNS = {"valid_from", "valid_to", "scraped_at"}

PRICE_NOTE = (
    "price: $0.0015 per deal, +$0.0005 per deal with --details, and from 2026-10-11 19:57 UTC a $0.001\n"
    "start fee per run (Apify API, read 2026-10-06; the Store page shows the current price).\n"
    "Not affiliated with Flipp or any retailer."
)


def load_dotenv(path: Path = Path(__file__).with_name(".env")) -> None:
    """Minimal .env reader so the scripts have no extra dependency."""
    if not path.exists():
        return
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def field(obj, camel: str, snake: str):
    """apify-client v1/v2 return dicts, v3 returns model objects. Support both."""
    if obj is None:
        return None
    return obj.get(camel) if isinstance(obj, dict) else getattr(obj, snake, None)


def is_postal(s: str) -> bool:
    """Canadian postal code (M5V 2T6, M5V2T6) or US ZIP (10001, 10001-1234). The first three
    characters of a postal code ("M5V") are not enough for Flipp."""
    s = s.strip()
    return bool(re.fullmatch(r"[A-Za-z]\d[A-Za-z][ -]?\d[A-Za-z]\d", s) or re.fullmatch(r"\d{5}(-?\d{4})?", s))


def split_list(v: str) -> list[str]:
    return [s.strip() for s in str(v).split(",") if s.strip()]


def read_list(path: str) -> list[str]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [s.strip() for s in lines if s.strip() and not s.strip().startswith("#")]


# $/kg is only meaningful for deals priced by weight (the Actor's README: "$/kg for weight-priced
# items"). A few in-store prices print an estimate such as "(est.) ea ($17.61/1kg ...)" and come back
# with no unit and a $/kg that doesn't match; for those price_per_kg is left empty.
WEIGHT_UNITS = {"lb", "kg", "100g"}


def _text(parts) -> str | None:
    s = " ".join(str(p).strip() for p in parts if p is not None and str(p).strip() != "")
    return s or None


def to_row(r: dict) -> dict:
    """One Actor row (or one row these scripts wrote earlier) -> one flat output row."""
    if "item_type" in r:
        return {c: r.get(c) for c in COLUMNS}
    mb = r.get("multiBuy") if isinstance(r.get("multiBuy"), dict) else None
    cats = [c for c in (r.get("flyerCategories") or []) if c] if isinstance(r.get("flyerCategories"), list) else []
    l1, l2 = r.get("categoryL1"), r.get("categoryL2")
    per_kg = r.get("pricePerKg") if r.get("priceUnit") in WEIGHT_UNITS else None
    each = r.get("effectivePrice")
    return {
        "search_term": r.get("searchTerm"),
        "rank": None,  # set by lowest_price_first()
        "price_basis": "kg" if per_kg is not None else ("item" if each is not None and float(each) > 0 else None),
        "name": r.get("name"),
        "brand": r.get("brand"),
        "merchant": r.get("merchant"),
        "effective_price": each,
        "price_per_kg": per_kg,
        "price_unit": r.get("priceUnit"),
        "price_as_printed": _text([r.get("prePriceText"), r.get("priceText"), r.get("postPriceText")]),
        "sale_story": r.get("saleStory"),
        "current_price": r.get("currentPrice"),
        "original_price": r.get("originalPrice"),
        "discount_pct": r.get("discountPct"),
        "dollars_off": r.get("dollarsOff"),
        "multi_buy_qty": mb.get("quantity") if mb else None,
        "multi_buy_total": mb.get("totalPrice") if mb else None,
        "currency": r.get("currency"),
        "valid_from": r.get("validFrom"),
        "valid_to": r.get("validTo"),
        "item_type": r.get("itemType"),
        "store_address": r.get("storeAddress"),
        "description": r.get("description"),
        "sku": r.get("sku"),
        "product_url": r.get("productUrl"),
        "image_url": r.get("imageUrl"),
        "flyer_id": r.get("flyerId"),
        "flyer_name": r.get("flyerName"),
        "flyer_categories": " | ".join(str(c) for c in cats) if cats else None,
        "product_category": f"{l1} > {l2}" if l1 and l2 else (l1 if l1 is not None else l2),
        "in_store_only": r.get("inStoreOnly"),
        "rating": r.get("rating"),
        "review_count": r.get("reviewCount"),
        "postal_code": r.get("postalCode"),
        "change_type": r.get("changeType"),
        "previous_price": r.get("previousPrice"),
        "item_id": r.get("itemId"),
        "merchant_id": r.get("merchantId"),
        "scraped_at": r.get("scrapedAt"),
    }


def is_deal(r) -> bool:
    return isinstance(r, dict) and ("item_type" in r or "itemType" in r)


def lowest_price_first(rows: list[dict]) -> list[dict]:
    """Lowest price first within each search term (terms keep the order they first appear in): deals
    priced by weight by $/kg, then deals priced per item by the per-item price, then deals with no
    single price ("20% off"). The two price bases aren't comparable, so each gets its own rank:
    1 = lowest for that search term and price basis."""
    order: dict[str, int] = {}
    for r in rows:
        order.setdefault(r.get("search_term") or "", len(order))
    basis = {"kg": 0, "item": 1}

    def key(r):
        b = r.get("price_basis")
        v = float(r["price_per_kg"]) if b == "kg" else float(r["effective_price"]) if b == "item" else 0.0
        return (order[r.get("search_term") or ""], basis.get(b, 2), v)

    ranked = sorted(rows, key=key)
    counter: dict[tuple, int] = {}
    for r in ranked:
        if not r.get("price_basis"):
            r["rank"] = None
            continue
        k = (r.get("search_term") or "", r["price_basis"])
        counter[k] = counter.get(k, 0) + 1
        r["rank"] = counter[k]
    return ranked


def csv_value(v):
    """Text that starts with = + - @, a tab or a carriage return gets a leading ' so Excel and Google
    Sheets show it as text instead of running it as a formula."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _parse_time(v):
    try:
        return datetime.fromisoformat(str(v).replace("Z", "+00:00"))
    except ValueError:
        return None


def write_xlsx(path: str, rows: list[dict]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("Writing .xlsx needs openpyxl: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Deals"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([None] * len(COLUMNS))
        line = ws.max_row
        for i, col in enumerate(COLUMNS, start=1):
            v = r.get(col)
            if v is None or v == "":
                continue
            fmt = None
            if col in NUMBER_COLUMNS:
                v = float(v) if isinstance(v, str) else v
            elif col in DATE_COLUMNS and _parse_time(v):
                # Excel has no time zones: the value is UTC.
                v = _parse_time(v).astimezone(timezone.utc).replace(tzinfo=None)
                fmt = "yyyy-mm-dd hh:mm"
            elif not isinstance(v, bool):
                v = ILLEGAL_CHARACTERS_RE.sub("", str(v))
            cell = ws.cell(row=line, column=i, value=v)
            if isinstance(v, str):
                cell.data_type = "s"  # never a formula, even if the text starts with =
            if fmt:
                cell.number_format = fmt
    widths = {"name": 45, "merchant": 18, "price_as_printed": 16, "sale_story": 18, "store_address": 36,
              "description": 40, "product_url": 40, "image_url": 40, "flyer_categories": 26,
              "product_category": 30, "valid_from": 17, "valid_to": 17, "scraped_at": 17}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 12)
    ws.freeze_panes = "A2"
    wb.save(path)


def write_rows(path: str, rows: list[dict]) -> None:
    low = path.lower()
    if low.endswith(".xlsx"):
        write_xlsx(path, rows)
        return
    p = Path(path)
    if low.endswith(".jsonl"):
        with p.open("w", encoding="utf-8") as f:
            for r in rows:
                f.write(json.dumps(r, ensure_ascii=False, separators=(",", ":")) + "\n")
        return
    # UTF-8 with a BOM, so Excel shows accents (Québec, crème) correctly when it opens the file.
    with p.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for r in rows:
            w.writerow([csv_value(r.get(c)) for c in COLUMNS])


def check_out(parser, path: str) -> None:
    if not re.search(r"\.(csv|xlsx|jsonl)$", path, re.I):
        parser.error(f"{path}: output files must end in .csv, .xlsx or .jsonl")


def read_from_file(path: str) -> list[dict]:
    """A JSON array, a single JSON object or JSONL: an Apify Console export of the Actor's rows,
    or a .jsonl file these scripts wrote."""
    text = Path(path).read_text(encoding="utf-8-sig").strip()
    if not text:
        return []
    data = json.loads(text) if text.startswith("[") else [json.loads(line) for line in text.splitlines() if line.strip()]
    return data if isinstance(data, list) else [data]


def start_fee_now() -> float:
    return START_FEE if datetime.now(timezone.utc) >= datetime.fromisoformat(START_FEE_FROM) else 0.0


def deals_for_cap(max_charge: float, details: bool) -> int:
    """How many deals a cost cap pays for."""
    per = PRICE_PER_DEAL + (PRICE_PER_DETAILS if details else 0)
    return max(0, int((max_charge - start_fee_now()) / per + 1e-9))


def apify_client():
    load_dotenv()
    token = os.environ.get("APIFY_TOKEN")
    if not token:
        sys.exit("Set APIFY_TOKEN (see .env.example), or use --from <file>")
    from apify_client import ApifyClient

    return ApifyClient(token)


def run_actor(client, run_input: dict, max_charge: float):
    """Start a run, wait for it, print how it ended. max_total_charge_usd is a hard cost cap Apify enforces."""
    run = client.actor(ACTOR).call(run_input=run_input, max_total_charge_usd=Decimal(str(max_charge)))
    if run is None:
        sys.exit("The run did not finish")
    status = field(run, "status", "status")
    status = getattr(status, "value", status)
    message = field(run, "statusMessage", "status_message")
    print(f"Run {field(run, 'id', 'id')} finished: {status}" + (f" ({message})" if message else ""), file=sys.stderr)
    return run, status


def dataset_items(client, dataset_id: str) -> list[dict]:
    return list(client.dataset(dataset_id).iterate_items())


def export_deals(args, run_input: dict, sort: bool) -> None:
    """Shared tail of search.py and flyers.py: new run, --dataset or --from -> rows -> file."""
    failed_run = False
    if args.from_file:
        items = read_from_file(args.from_file)
    else:
        client = apify_client()
        dataset_id = args.dataset
        if not dataset_id:
            run, status = run_actor(client, run_input, args.max_charge)
            failed_run = status != "SUCCEEDED"
            dataset_id = field(run, "defaultDatasetId", "default_dataset_id")
        items = dataset_items(client, dataset_id)
    if failed_run and not items:
        sys.exit(f"No deals, so {args.out} was left as it was.")
    rows = [to_row(it) for it in items if is_deal(it)]
    if sort:
        rows = lowest_price_first(rows)
    write_rows(args.out, rows)
    print(f"Wrote {len(rows)} deals to {args.out}", file=sys.stderr)
